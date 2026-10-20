import pytest
from docx import Document

from app.core.utils import split_articles, tokenize_query, best_sentence, extract_text_with_config, OCRExtractionError


def test_split_articles():
    text = "第一条 内容1\n第二条 内容2"
    items = split_articles(text)
    assert len(items) == 2
    assert items[0] == ("第一条", "内容1")
    assert items[1] == ("第二条", "内容2")


def test_tokenize_query():
    q = "法律法规 search"
    tokens = tokenize_query(q)
    assert "法律法规" in tokens
    assert "search" in tokens


def test_best_sentence():
    text = "句子一。句子二包含关键词。"
    tokens = ["关键词"]
    s, score = best_sentence(text, tokens)
    assert s == "句子二包含关键词"
    assert score == 1


def test_tokenize_query_tax_terms():
    q = "税收征收管理法实施细则 增值税专用缴款书"
    tokens = tokenize_query(q)
    assert "税收征收管理法实施细则" in tokens
    assert "增值税专用缴款书" in tokens


def test_extract_text_with_config_raises_when_ocr_returns_empty(monkeypatch):
    monkeypatch.setattr(
        "app.services.pdf_pipeline.extract_pdf_text_with_pipeline",
        lambda *_args, **_kwargs: ("", {"ocr_used": True, "ocr_engine": "mineru", "page_count": 1}),
    )

    cfg = {"ocr_enabled": True, "ocr_min_text_length": 10, "ocr_languages": "chi_sim+eng", "ocr_dpi": 220}
    with pytest.raises(OCRExtractionError):
        extract_text_with_config(cfg, "dummy.pdf")


def test_extract_text_with_config_raises_when_ocr_engine_failed(monkeypatch):
    def _boom(*_args, **_kwargs):
        raise RuntimeError("No module named 'torchvision'")

    monkeypatch.setattr("app.services.pdf_pipeline.extract_pdf_text_with_pipeline", _boom)

    cfg = {"ocr_enabled": True, "ocr_min_text_length": 10, "ocr_languages": "chi_sim+eng", "ocr_dpi": 220}
    with pytest.raises(OCRExtractionError):
        extract_text_with_config(cfg, "dummy.pdf")


def test_extract_text_with_config_reads_docx_paragraphs_and_tables(tmp_path):
    file_path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("Main Clause")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Deposit"
    table.cell(1, 1).text = "100"
    doc.save(file_path)

    text, meta = extract_text_with_config({"ocr_enabled": False}, str(file_path))

    assert meta["ext"] == ".docx"
    assert "Main Clause" in text
    assert "Item\tAmount" in text
    assert "Deposit\t100" in text
