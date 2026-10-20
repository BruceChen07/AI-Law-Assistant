import json
import logging
import os
import re
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.core.mineru_ocr import probe_mineru_capabilities, run_mineru_extract
from app.core.utils import extract_text_with_config
from app.services.crud import (
    create_document_digitization_record,
    get_latest_successful_document_digitization_by_document,
    get_next_document_digitization_version,
)


logger = logging.getLogger("law_assistant")

_HEADING_LEVEL_TO_STYLE = {
    1: "Title",
    2: "Heading 1",
    3: "Heading 2",
    4: "Heading 3",
}


def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _safe_int(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except Exception:
        return int(default)


def _safe_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except Exception:
        return float(default)


def _normalize_text(value: Any) -> str:
    text = str(value or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return "\n".join([line.rstrip() for line in text.split("\n")]).strip()


def _load_json(path: str) -> Any:
    if not path or not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _looks_like_heading(text: str) -> bool:
    s = str(text or "").strip()
    if not s:
        return False
    if len(s) <= 28 and re.match(r"^(第[一二三四五六七八九十百千0-9]+[章节条款]|[0-9]+(?:\.[0-9]+){0,3}|[一二三四五六七八九十]+、|(?:Article|Section|Chapter|Part)\s+[0-9IVXLCM]+(?:\.[0-9]+)*)", s, flags=re.IGNORECASE):
        return True
    return len(s) <= 18 and not any(ch in s for ch in "，。；;,.")


def _sort_page_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    def _key(row: Dict[str, Any]) -> Tuple[int, float, float]:
        bbox = row.get("bbox") if isinstance(row.get("bbox"), list) else []
        left = _safe_float(bbox[0], 0.0) if len(bbox) >= 1 else 0.0
        top = _safe_float(bbox[1], 0.0) if len(bbox) >= 2 else 0.0
        return (_safe_int(row.get("_index"), 0), round(top, 3), round(left, 3))

    return sorted(rows, key=_key)


def _page_size_from_middle(middle: Any, page_no: int) -> Tuple[float, float]:
    pdf_info = middle.get("pdf_info") if isinstance(middle, dict) else None
    if not isinstance(pdf_info, list):
        return 595.0, 842.0
    idx = max(0, int(page_no) - 1)
    if idx >= len(pdf_info):
        return 595.0, 842.0
    page = pdf_info[idx]
    page_size = page.get("page_size") if isinstance(page, dict) else None
    if isinstance(page_size, list) and len(page_size) >= 2:
        return max(1.0, _safe_float(page_size[0], 595.0)), max(1.0, _safe_float(page_size[1], 842.0))
    return 595.0, 842.0


def _resolve_asset_path(base_dir: str, rel_path: str) -> str:
    raw = str(rel_path or "").strip()
    if not raw:
        return ""
    if os.path.isabs(raw):
        return raw
    return os.path.abspath(os.path.join(base_dir, raw.replace("/", os.sep)))


def _paragraph_indent_inches(left: float, page_width: float) -> float:
    if page_width <= 0:
        return 0.0
    ratio = max(0.0, min(left / page_width, 1.0))
    return round(min(1.4, ratio * 1.8), 2)


def _paragraph_alignment(left: float, width: float, page_width: float) -> WD_ALIGN_PARAGRAPH:
    if page_width <= 0:
        return WD_ALIGN_PARAGRAPH.LEFT
    center = left + (width / 2.0)
    center_ratio = center / page_width
    width_ratio = width / page_width
    if 0.38 <= center_ratio <= 0.62 and width_ratio <= 0.55:
        return WD_ALIGN_PARAGRAPH.CENTER
    return WD_ALIGN_PARAGRAPH.LEFT


def _set_paragraph_geometry(paragraph, bbox: List[float], page_width: float) -> None:
    left = _safe_float(bbox[0], 0.0) if len(bbox) >= 1 else 0.0
    width = _safe_float(bbox[2], 0.0) if len(bbox) >= 3 else 0.0
    paragraph.paragraph_format.left_indent = Inches(_paragraph_indent_inches(left, page_width))
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.alignment = _paragraph_alignment(left, width, page_width)


def _parse_table_html(table_html: str) -> List[List[Dict[str, Any]]]:
    raw = str(table_html or "").strip()
    if not raw:
        return []
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        try:
            root = ET.fromstring(f"<root>{raw}</root>")
        except ET.ParseError:
            return []
    table = root if root.tag.lower().endswith("table") else root.find(".//table")
    if table is None:
        return []

    rows: List[List[Dict[str, Any]]] = []
    occupied: Dict[Tuple[int, int], bool] = {}
    max_cols = 0

    for row_idx, tr in enumerate(table.findall(".//tr")):
        current_row: List[Dict[str, Any]] = []
        col_idx = 0
        for cell in list(tr):
            tag = str(cell.tag or "").lower().split("}")[-1]
            if tag not in {"td", "th"}:
                continue
            while occupied.get((row_idx, col_idx)):
                col_idx += 1
            rowspan = max(1, _safe_int(cell.attrib.get("rowspan"), 1))
            colspan = max(1, _safe_int(cell.attrib.get("colspan"), 1))
            text = _normalize_text(" ".join(cell.itertext()))
            entry = {
                "col_idx": col_idx,
                "rowspan": rowspan,
                "colspan": colspan,
                "text": text,
            }
            current_row.append(entry)
            for rr in range(row_idx, row_idx + rowspan):
                for cc in range(col_idx, col_idx + colspan):
                    if rr == row_idx and cc == col_idx:
                        continue
                    occupied[(rr, cc)] = True
            col_idx += colspan
        max_cols = max(max_cols, col_idx)
        rows.append(current_row)

    matrix: List[List[Dict[str, Any]]] = []
    for _ in rows:
        matrix.append([{"text": "", "anchor": False, "rowspan": 1, "colspan": 1} for _ in range(max_cols)])
    for row_idx, row in enumerate(rows):
        for entry in row:
            col_idx = int(entry["col_idx"])
            matrix[row_idx][col_idx] = {
                "text": str(entry["text"] or ""),
                "anchor": True,
                "rowspan": int(entry["rowspan"]),
                "colspan": int(entry["colspan"]),
            }
    return matrix


def _apply_table_matrix(doc: Document, matrix: List[List[Dict[str, Any]]]) -> None:
    if not matrix or not matrix[0]:
        return
    row_count = len(matrix)
    col_count = len(matrix[0])
    table = doc.add_table(rows=row_count, cols=col_count)
    table.style = "Table Grid"
    for row_idx, row in enumerate(matrix):
        for col_idx, cell_info in enumerate(row):
            if not cell_info.get("anchor"):
                continue
            anchor = table.cell(row_idx, col_idx)
            anchor.text = str(cell_info.get("text") or "")
            rowspan = max(1, _safe_int(cell_info.get("rowspan"), 1))
            colspan = max(1, _safe_int(cell_info.get("colspan"), 1))
            if rowspan > 1 or colspan > 1:
                end_row = min(row_count - 1, row_idx + rowspan - 1)
                end_col = min(col_count - 1, col_idx + colspan - 1)
                if end_row != row_idx or end_col != col_idx:
                    anchor.merge(table.cell(end_row, end_col))


def _add_text_block(doc: Document, row: Dict[str, Any], page_width: float) -> None:
    text = _normalize_text(row.get("text"))
    if not text:
        return
    level = _safe_int(row.get("text_level"), 0)
    style_name = _HEADING_LEVEL_TO_STYLE.get(level)
    if style_name is None and _looks_like_heading(text):
        style_name = "Heading 2"
    paragraph = doc.add_paragraph(style=style_name) if style_name else doc.add_paragraph()
    paragraph.add_run(text)
    bbox = row.get("bbox") if isinstance(row.get("bbox"), list) else []
    _set_paragraph_geometry(paragraph, bbox, page_width)


def _add_table_block(doc: Document, row: Dict[str, Any], page_width: float, page_images: int) -> int:
    caption_items = row.get("table_caption") if isinstance(row.get("table_caption"), list) else []
    for item in caption_items:
        caption = _normalize_text(item)
        if caption:
            p = doc.add_paragraph(caption, style="Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    matrix = _parse_table_html(str(row.get("table_body") or ""))
    _apply_table_matrix(doc, matrix)
    img_path = _resolve_asset_path(os.path.dirname(str(row.get("_content_list_path") or "")), str(row.get("img_path") or ""))
    if img_path and os.path.exists(img_path):
        bbox = row.get("bbox") if isinstance(row.get("bbox"), list) else []
        width = _safe_float(bbox[2], 0.0) - _safe_float(bbox[0], 0.0)
        width_in = min(6.5, max(1.4, (width / max(1.0, page_width)) * 7.0))
        try:
            doc.add_picture(img_path, width=Inches(width_in))
            page_images += 1
        except Exception:
            logger.warning("digitize_table_image_add_failed file=%s", img_path, exc_info=True)
    footnotes = row.get("table_footnote") if isinstance(row.get("table_footnote"), list) else []
    for item in footnotes:
        note = _normalize_text(item)
        if note:
            doc.add_paragraph(note)
    return page_images


def _add_image_block(doc: Document, row: Dict[str, Any], page_width: float, page_images: int) -> int:
    img_path = _resolve_asset_path(os.path.dirname(str(row.get("_content_list_path") or "")), str(row.get("img_path") or ""))
    if not img_path or not os.path.exists(img_path):
        return page_images
    bbox = row.get("bbox") if isinstance(row.get("bbox"), list) else []
    width = _safe_float(bbox[2], 0.0) - _safe_float(bbox[0], 0.0)
    width_in = min(6.5, max(1.4, (width / max(1.0, page_width)) * 7.0))
    try:
        doc.add_picture(img_path, width=Inches(width_in))
        page_images += 1
    except Exception:
        logger.warning("digitize_image_add_failed file=%s", img_path, exc_info=True)
    caption_items = row.get("image_caption") if isinstance(row.get("image_caption"), list) else []
    for item in caption_items:
        caption = _normalize_text(item)
        if caption:
            p = doc.add_paragraph(caption, style="Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footnotes = row.get("image_footnote") if isinstance(row.get("image_footnote"), list) else []
    for item in footnotes:
        note = _normalize_text(item)
        if note:
            doc.add_paragraph(note)
    return page_images


def rebuild_docx_from_mineru_artifacts(
    *,
    content_list_path: str,
    middle_path: str,
    output_docx_path: str,
    title: str = "",
) -> Dict[str, Any]:
    content_list = _load_json(content_list_path)
    middle = _load_json(middle_path)
    rows = content_list if isinstance(content_list, list) else []
    if not rows:
        raise RuntimeError("digitize_content_list_empty")

    grouped: Dict[int, List[Dict[str, Any]]] = {}
    for index, raw in enumerate(rows, start=1):
        if not isinstance(raw, dict):
            continue
        page_idx = _safe_int(raw.get("page_idx"), -1)
        if page_idx < 0:
            continue
        row = dict(raw)
        row["_index"] = index
        row["_content_list_path"] = content_list_path
        grouped.setdefault(page_idx + 1, []).append(row)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    if title:
        title_paragraph = doc.add_paragraph(style="Title")
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_paragraph.add_run(_normalize_text(title))

    paragraph_count = 0
    table_count = 0
    image_count = 0
    page_total = len(grouped)

    for page_no in sorted(grouped.keys()):
        page_width, _page_height = _page_size_from_middle(middle, page_no)
        page_rows = _sort_page_rows(grouped.get(page_no) or [])
        for row in page_rows:
            block_type = str(row.get("type") or "text").strip().lower()
            if block_type in {"discarded"}:
                continue
            if block_type == "text":
                before = len(doc.paragraphs)
                _add_text_block(doc, row, page_width)
                paragraph_count += max(0, len(doc.paragraphs) - before)
                continue
            if block_type == "table":
                image_count = _add_table_block(doc, row, page_width, image_count)
                table_count += 1
                continue
            if block_type == "image":
                image_count = _add_image_block(doc, row, page_width, image_count)
                continue
            if block_type in {"title"}:
                before = len(doc.paragraphs)
                row = dict(row)
                row["text_level"] = row.get("text_level") or 1
                _add_text_block(doc, row, page_width)
                paragraph_count += max(0, len(doc.paragraphs) - before)
                continue
            text = _normalize_text(row.get("text"))
            if text:
                before = len(doc.paragraphs)
                _add_text_block(doc, row, page_width)
                paragraph_count += max(0, len(doc.paragraphs) - before)
        if page_no < page_total:
            doc.add_section(WD_SECTION.NEW_PAGE)

    os.makedirs(os.path.dirname(output_docx_path), exist_ok=True)
    doc.save(output_docx_path)
    return {
        "paragraph_count": paragraph_count,
        "table_count": table_count,
        "image_count": image_count,
        "page_count": page_total,
    }


def _text_match_ratio(a: str, b: str) -> float:
    left = re.sub(r"\s+", "", str(a or ""))
    right = re.sub(r"\s+", "", str(b or ""))
    if not left or not right:
        return 0.0
    return round(SequenceMatcher(None, left[:120000], right[:120000]).ratio(), 4)


def _digitization_storage_dir(cfg: Dict[str, Any], document_id: str, version_no: int) -> str:
    files_dir = str(cfg.get("files_dir") or "")
    root = os.path.join(files_dir, "contract_digitization", str(document_id), f"v{int(version_no)}")
    os.makedirs(root, exist_ok=True)
    return root


def _digitized_docx_name(source_file_path: str) -> str:
    stem = Path(str(source_file_path or "contract")).stem
    return f"{stem}_digitized.docx"


def ensure_digitized_docx(
    cfg: Dict[str, Any],
    *,
    document_id: str,
    source_file_path: str,
    original_filename: str = "",
) -> Dict[str, Any]:
    ext = str(Path(source_file_path).suffix or "").lower()
    if ext != ".pdf":
        return {
            "status": "skipped",
            "document_id": document_id,
            "source_file_path": source_file_path,
            "source_ext": ext,
            "digitized_docx_path": source_file_path,
            "digitization_used": False,
            "layout_retained": ext == ".docx",
            "version_no": 0,
        }

    cached = get_latest_successful_document_digitization_by_document(cfg, document_id)
    if cached:
        cached_docx = str(cached.get("digitized_docx_path") or "")
        cached_source = str(cached.get("source_file_path") or "")
        if cached_docx and os.path.exists(cached_docx) and cached_source == source_file_path:
            out = dict(cached)
            out["digitization_used"] = True
            out["reused"] = True
            return out

    capabilities = probe_mineru_capabilities(force_refresh=True)
    if not capabilities.get("mineru_cli_found"):
        raise RuntimeError("digitize_ocr_unavailable: MinerU CLI not found")

    version_no = get_next_document_digitization_version(cfg, document_id)
    output_dir = _digitization_storage_dir(cfg, document_id, version_no)
    docx_path = os.path.join(output_dir, _digitized_docx_name(source_file_path))
    mineru_dir = os.path.join(output_dir, "mineru")
    os.makedirs(mineru_dir, exist_ok=True)

    logger.info(
        "digitize_pdf_start document_id=%s version=%s file=%s output_dir=%s",
        document_id,
        version_no,
        source_file_path,
        output_dir,
    )
    try:
        result = run_mineru_extract(source_file_path, output_dir=mineru_dir)
        content_list_path = str(result.get("content_list_path") or "")
        middle_path = str(result.get("middle_path") or "")
        md_path = str(result.get("md_path") or "")
        if not content_list_path or not os.path.exists(content_list_path):
            raise RuntimeError("digitize_content_list_missing")

        reconstructed = rebuild_docx_from_mineru_artifacts(
            content_list_path=content_list_path,
            middle_path=middle_path,
            output_docx_path=docx_path,
            title=Path(original_filename or source_file_path).stem,
        )
        docx_text, _docx_meta = extract_text_with_config(cfg, docx_path)
        source_text = str(result.get("text") or "")
        text_ratio = _text_match_ratio(source_text, docx_text)
        record = {
            "id": f"{document_id}::digitize::{version_no}",
            "document_id": document_id,
            "version_no": version_no,
            "status": "done",
            "source_file_path": source_file_path,
            "source_ext": ext,
            "digitized_docx_path": docx_path,
            "artifact_dir": output_dir,
            "md_path": md_path,
            "middle_path": middle_path,
            "content_list_path": content_list_path,
            "ocr_engine": "mineru",
            "ocr_used": 1,
            "layout_retained": 1,
            "source_text_length": len(source_text),
            "digitized_text_length": len(docx_text),
            "text_match_ratio": text_ratio,
            "paragraph_count": int(reconstructed.get("paragraph_count") or 0),
            "table_count": int(reconstructed.get("table_count") or 0),
            "image_count": int(reconstructed.get("image_count") or 0),
            "page_count": int(reconstructed.get("page_count") or 0),
            "note": "",
            "created_at": _utc_now_iso(),
            "updated_at": _utc_now_iso(),
        }
        create_document_digitization_record(cfg, record)
        logger.info(
            "digitize_pdf_done document_id=%s version=%s docx=%s text_ratio=%s paragraphs=%s tables=%s images=%s",
            document_id,
            version_no,
            docx_path,
            text_ratio,
            record["paragraph_count"],
            record["table_count"],
            record["image_count"],
        )
        record["digitization_used"] = True
        record["reused"] = False
        return record
    except Exception as exc:
        create_document_digitization_record(
            cfg,
            {
                "id": f"{document_id}::digitize::{version_no}",
                "document_id": document_id,
                "version_no": version_no,
                "status": "failed",
                "source_file_path": source_file_path,
                "source_ext": ext,
                "digitized_docx_path": "",
                "artifact_dir": output_dir,
                "md_path": "",
                "middle_path": "",
                "content_list_path": "",
                "ocr_engine": "mineru",
                "ocr_used": 1,
                "layout_retained": 0,
                "source_text_length": 0,
                "digitized_text_length": 0,
                "text_match_ratio": 0,
                "paragraph_count": 0,
                "table_count": 0,
                "image_count": 0,
                "page_count": 0,
                "note": str(exc),
                "created_at": _utc_now_iso(),
                "updated_at": _utc_now_iso(),
            },
        )
        raise
